#!/usr/bin/env python3
"""Generate Facebook GROUP report – comparing multiple pages within a group."""
import openpyxl
import json
import os
import re
import csv
import math
from datetime import datetime
from collections import defaultdict, Counter

DIR = os.path.dirname(os.path.abspath(__file__))
XLSX = '/Users/arturs25/Downloads/facebook_posts-2026-02-24.xlsx'

# ── GROUP CONFIGS ──
import sys

GROUPS = {
    'parvalde': {
        'name': 'Pārvalde un Tūrisms',
        'slug': 'parvalde_turisms',
        'pages': ['Ventspils', 'VisitVentspils'],
        'short': {'Ventspils': 'Ventspils', 'VisitVentspils': 'VisitVentspils'},
        'research': '/Users/arturs25/Downloads/Facebook Research - Ventspils, VisitVentspils/',
    },
    'izglitiba': {
        'name': 'Izglītība',
        'slug': 'izglitiba',
        'pages': [
            'Ventspils Izglītības pārvalde',
            'Ventspils Jaunrades nams',
            'Sporta skola "Spars"',
            'Ventspils Jauniešu portāls',
            'Ventspils Digitālais centrs',
            'Zinātnes centrs "Vizium"',
        ],
        'short': {
            'Ventspils Izglītības pārvalde': 'Izgl. pārvalde',
            'Ventspils Jaunrades nams': 'Jaunrades nams',
            'Sporta skola "Spars"': 'Spars',
            'Ventspils Jauniešu portāls': 'Jauniešu portāls',
            'Ventspils Digitālais centrs': 'Digitālais centrs',
            'Zinātnes centrs "Vizium"': 'Vizium',
        },
        'research': '/Users/arturs25/Downloads/Facebook Research - Sporta skola _Spars_, Ventspils Digitālais centrs, Ventspils Izglītības pārvalde, Ventspils Jauniešu portāls, Ventspils Jaunrades nams, Zinātnes centrs _Vizium_/',
    },
    'socialais': {
        'name': 'Sociālie pakalpojumi un Veselības aprūpe',
        'slug': 'socialais_veseliba',
        'pages': [
            'Ventspils Sociālais dienests',
            'Ziemeļkurzemes reģionālā slimnīca',
            'Ventspils poliklīnika',
        ],
        'short': {
            'Ventspils Sociālais dienests': 'Soc. dienests',
            'Ziemeļkurzemes reģionālā slimnīca': 'Slimnīca',
            'Ventspils poliklīnika': 'Poliklīnika',
        },
        'research': '/Users/arturs25/Downloads/Facebook Research - Ventspils poliklīnika, Ventspils Sociālais dienests, Ziemeļkurzemes reģionālā slimnīca/',
    },
    'sports': {
        'name': 'Sports',
        'slug': 'sports',
        'pages': [
            'Olimpiskais centrs Ventspils',
            'Ventspils Sporta pārvalde',
            'Piedzīvojumu Parks / Lemberga hūte',
            'Ventspils Ūdens piedzīvojumu parks',
        ],
        'short': {
            'Olimpiskais centrs Ventspils': 'OC Ventspils',
            'Ventspils Sporta pārvalde': 'Sporta pārvalde',
            'Piedzīvojumu Parks / Lemberga hūte': 'Piedzīv. parks',
            'Ventspils Ūdens piedzīvojumu parks': 'Ūdens parks',
        },
        'research': '/Users/arturs25/Downloads/Facebook Research - Olimpiskais centrs Ventspils, Piedzīvojumu Parks _ Lemberga hūte, Ventspils Sporta pārvalde, Ventspils Ūdens piedzīvojumu parks/',
    },
    'kultura': {
        'name': 'Kultūra',
        'slug': 'kultura',
        'pages': [
            'Ventspils Kultūras centrs',
            'Teātra nams JŪRAS VĀRTI',
            'Koncertzāle Latvija',
            'Ventspils bibliotēka',
            'Ventspils muzejs',
        ],
        'short': {
            'Ventspils Kultūras centrs': 'Kultūras centrs',
            'Teātra nams JŪRAS VĀRTI': 'Jūras vārti',
            'Koncertzāle Latvija': 'Konc. Latvija',
            'Ventspils bibliotēka': 'Bibliotēka',
            'Ventspils muzejs': 'Muzejs',
        },
        'research': '/Users/arturs25/Downloads/Facebook Research - Koncertzāle Latvija, Teātra nams JŪRAS VĀRTI, Ventspils bibliotēka, Ventspils Kultūras centrs, Ventspils muzejs/',
    },
    'komunalie': {
        'name': 'Komunālie pakalpojumi',
        'slug': 'komunalie',
        'pages': [
            'Pašvaldības SIA ŪDEKA',
            'Ventspils labiekārtošanas kombināts',
            'Ventspils Siltums',
            'Ventspils nekustamie īpašumi',
        ],
        'short': {
            'Pašvaldības SIA ŪDEKA': 'Ūdeka',
            'Ventspils labiekārtošanas kombināts': 'Labiekārt. komb.',
            'Ventspils Siltums': 'Ventspils Siltums',
            'Ventspils nekustamie īpašumi': 'Nekust. īpašumi',
        },
        'research': '/Users/arturs25/Downloads/Facebook Research - Pašvaldības SIA ŪDEKA, Ventspils labiekārtošanas kombināts, Ventspils nekustamie īpašumi, Ventspils Siltums/',
    },
    'citi': {
        'name': 'Citi',
        'slug': 'citi',
        'pages': [
            'Ventspils Komunālā pārvalde',
            'Pašvaldības SIA Ventspils Reiss',
            'Ventspils Pašvaldības policija',
            'Ventspils tirgus',
        ],
        'short': {
            'Ventspils Komunālā pārvalde': 'Komunālā pārvalde',
            'Pašvaldības SIA Ventspils Reiss': 'Ventspils Reiss',
            'Ventspils Pašvaldības policija': 'Pašv. policija',
            'Ventspils tirgus': 'Ventspils tirgus',
        },
        'research': '/Users/arturs25/Downloads/Facebook Research - Pašvaldības SIA Ventspils Reiss, Ventspils Komunālā pārvalde, Ventspils Pašvaldības policija, Ventspils tirgus/',
    },
    'partneri': {
        'name': 'Sadarbības partneri',
        'slug': 'partneri',
        'pages': [
            'Ventspils Augstskola',
            'Ventspils Tehnikums',
            'Ventspils Mūzikas vidusskola',
        ],
        'short': {
            'Ventspils Augstskola': 'Augstskola',
            'Ventspils Tehnikums': 'Tehnikums',
            'Ventspils Mūzikas vidusskola': 'Mūzikas vidussk.',
        },
        'research': '/Users/arturs25/Downloads/Facebook Research - Ventspils Augstskola, Ventspils Mūzikas vidusskola, Ventspils Tehnikums/',
    },
}

# Color palette for up to 8 pages
ALL_COLORS = ['#1877F2', '#34D399', '#FBBF24', '#F87171', '#A78BFA', '#06b6d4', '#ec4899', '#f97316']

# Ordered list for prev/next navigation between group reports
GROUP_ORDER = ['parvalde', 'izglitiba', 'socialais', 'sports', 'kultura', 'komunalie', 'citi', 'partneri']

GROUP_KEY = sys.argv[1] if len(sys.argv) > 1 else 'parvalde'
CFG = GROUPS[GROUP_KEY]
GROUP_NAME = CFG['name']
GROUP_SLUG = CFG['slug']
PAGES = CFG['pages']
PAGE_SHORT = CFG['short']
PAGE_COLORS = {p: ALL_COLORS[i % len(ALL_COLORS)] for i, p in enumerate(PAGES)}
RESEARCH_DIR = CFG['research']

MONTH_NAMES = ['Jan','Feb','Mar','Apr','Mai','Jun','Jul','Aug','Sep','Okt','Nov','Dec']
DAY_NAMES = ['Pirmdiena','Otrdiena','Trešdiena','Ceturtdiena','Piektdiena','Sestdiena','Svētdiena']

def safe_int(v):
    if v is None: return 0
    try: return int(v)
    except: return 0

def safe_float(v):
    if v is None: return 0.0
    try: return float(v)
    except: return 0.0

def read_csv(filename):
    path = os.path.join(RESEARCH_DIR, filename)
    if not os.path.exists(path):
        return []
    with open(path, 'r', encoding='utf-8') as f:
        return list(csv.reader(f))


def load_all_data():
    print(f"Loading data for group: {GROUP_NAME}")
    print(f"  Pages: {', '.join(PAGES)}")
    wb = openpyxl.load_workbook(XLSX, read_only=True)
    ws = wb['Worksheet']

    page_posts = {p: [] for p in PAGES}

    for row in ws.iter_rows(min_row=2, values_only=True):
        pname = row[1]
        if pname not in PAGES:
            continue

        ts = row[12]
        hour = None
        if ts and isinstance(ts, (int, float)):
            try:
                hour = datetime.fromtimestamp(ts).hour
            except:
                pass

        teksts = str(row[3]) if row[3] else ''
        t_lower = teksts.lower()
        is_contest = any(k in t_lower for k in ['konkurs','laimē','izlozē','loterij','izloze','laimēt','piedalies un laimē','piedalies'])

        post = {
            'id': row[0],
            'page': pname,
            'type': row[2] or 'Other',
            'teksts': teksts,
            't_zimes': safe_int(row[4]),
            'picture': str(row[5]) if row[5] else '',
            'gif': str(row[6]).upper() == 'Y' if row[6] else False,
            'a_zimes': safe_int(row[7]),
            'a_prop': str(row[8]) if row[8] else '',
            'a_krasa': str(row[9]) if row[9] else '',
            'date': str(row[11]) if row[11] else '',
            'hour': hour,
            'day': safe_int(row[13]),
            'shares': safe_int(row[14]),
            'likes': safe_int(row[15]),
            'comments': safe_int(row[16]),
            'slc': safe_int(row[17]),
            'video_len': safe_float(row[18]),
            'views': safe_int(row[19]),
            'is_contest': is_contest,
            'fb_url': str(row[27]) if row[27] else '',
        }
        page_posts[pname].append(post)

    wb.close()
    for p in PAGES:
        print(f"  {p}: {len(page_posts[p])} posts")
    return page_posts


def find_csv(prefix):
    """Find a CSV file in research dir starting with given prefix."""
    import unicodedata
    if not os.path.isdir(RESEARCH_DIR):
        return None
    nfc_prefix = unicodedata.normalize('NFC', prefix)
    for f in os.listdir(RESEARCH_DIR):
        nfc_f = unicodedata.normalize('NFC', f)
        if nfc_f.startswith(nfc_prefix) and nfc_f.endswith('.csv'):
            return f  # Return original filename for os.path operations
    return None


def load_research():
    """Load supplementary data from research CSVs."""
    R = {}

    # Komunikācija
    fname = find_csv('Komunikācija')
    if fname:
        rows = read_csv(fname)
        if len(rows) > 1:
            headers = rows[0]
            R['komunikacija'] = {}
            for row in rows[1:]:
                key = row[0].strip().replace('"', '')
                R['komunikacija'][key] = {}
                for i, p in enumerate(headers[1:], 1):
                    R['komunikacija'][key][p.strip()] = row[i].strip() if i < len(row) else ''

    # Publikācijas
    fname = find_csv('Publikācijas')
    if fname:
        rows = read_csv(fname)
        if len(rows) > 1:
            headers = rows[0]
            R['publikacijas'] = {}
            for row in rows[1:]:
                key = row[0].strip()
                R['publikacijas'][key] = {}
                for i, p in enumerate(headers[1:], 1):
                    R['publikacijas'][key][p.strip()] = row[i].strip() if i < len(row) else ''

    # Populārākās saites
    fname = find_csv('Populārākās saites')
    if fname:
        rows = read_csv(fname)
        if len(rows) > 1:
            R['saites'] = []
            for row in rows[1:]:
                if len(row) >= 3:
                    R['saites'].append({
                        'domain': row[0].strip(),
                        'count': safe_int(row[1]),
                        'pct': row[2].strip() if len(row) > 2 else '',
                    })

    return R


def compute_page(posts, page_name):
    """Compute analytics for a single page (reused logic)."""
    D = {}
    n = len(posts)
    if n == 0:
        return D

    total_likes = sum(p['likes'] for p in posts)
    total_shares = sum(p['shares'] for p in posts)
    total_comments = sum(p['comments'] for p in posts)
    total_eng = total_likes + total_shares + total_comments
    avg_eng = round(total_eng / n, 1) if n else 0

    sorted_eng = sorted(p['slc'] for p in posts)
    median_eng = sorted_eng[n // 2] if n else 0

    D['kpi'] = {
        'page': page_name,
        'total_posts': n,
        'total_likes': total_likes,
        'total_shares': total_shares,
        'total_comments': total_comments,
        'total_engagement': total_eng,
        'avg_eng': avg_eng,
        'median_eng': median_eng,
    }

    viral_threshold = avg_eng * 3
    viral = [p for p in posts if p['slc'] >= viral_threshold]
    D['kpi']['viral_count'] = len(viral)
    D['kpi']['viral_pct'] = round(len(viral) / n * 100, 1)

    # Monthly
    monthly = defaultdict(lambda: {'posts': 0, 'eng': 0})
    for p in posts:
        d = p['date']
        if d and len(d) >= 7:
            try:
                m = int(d[5:7])
                monthly[m]['posts'] += 1
                monthly[m]['eng'] += p['slc']
            except:
                pass

    D['monthly'] = []
    for m in range(1, 13):
        md = monthly[m]
        D['monthly'].append({
            'month': m,
            'name': MONTH_NAMES[m - 1],
            'posts': md['posts'],
            'engagement': md['eng'],
            'avg_eng': round(md['eng'] / md['posts'], 1) if md['posts'] else 0,
        })

    # H1 vs H2
    h1 = [m for m in D['monthly'] if m['month'] <= 6]
    h2 = [m for m in D['monthly'] if m['month'] > 6]
    h1_posts = sum(m['posts'] for m in h1)
    h2_posts = sum(m['posts'] for m in h2)
    h1_eng = sum(m['engagement'] for m in h1)
    h2_eng = sum(m['engagement'] for m in h2)
    h1_avg = round(h1_eng / h1_posts, 1) if h1_posts else 0
    h2_avg = round(h2_eng / h2_posts, 1) if h2_posts else 0
    D['half_year'] = [
        {'name': 'H1', 'posts': h1_posts, 'avg_eng': h1_avg},
        {'name': 'H2', 'posts': h2_posts, 'avg_eng': h2_avg},
    ]

    # Weekday
    weekday_data = defaultdict(lambda: {'posts': 0, 'eng': 0})
    for p in posts:
        d = p['day']
        if 1 <= d <= 7:
            weekday_data[d]['posts'] += 1
            weekday_data[d]['eng'] += p['slc']

    D['weekday'] = []
    for d in range(1, 8):
        wd = weekday_data[d]
        D['weekday'].append({
            'day': d,
            'name': DAY_NAMES[d - 1],
            'posts': wd['posts'],
            'avg_eng': round(wd['eng'] / wd['posts'], 1) if wd['posts'] else 0,
        })

    # Work vs Weekend
    work = [p for p in posts if 1 <= p['day'] <= 5]
    weekend = [p for p in posts if p['day'] in (6, 7)]
    D['work_weekend'] = {
        'work_posts': len(work),
        'work_pct': round(len(work) / n * 100, 1),
        'work_avg': round(sum(p['slc'] for p in work) / len(work), 1) if work else 0,
        'weekend_posts': len(weekend),
        'weekend_pct': round(len(weekend) / n * 100, 1),
        'weekend_avg': round(sum(p['slc'] for p in weekend) / len(weekend), 1) if weekend else 0,
    }

    # Hourly
    hourly = defaultdict(lambda: {'posts': 0, 'eng': 0})
    for p in posts:
        if p['hour'] is not None:
            hourly[p['hour']]['posts'] += 1
            hourly[p['hour']]['eng'] += p['slc']

    D['hourly'] = []
    for h in range(24):
        hd = hourly[h]
        D['hourly'].append({
            'hour': f"{h:02d}:00",
            'posts': hd['posts'],
            'avg_eng': round(hd['eng'] / hd['posts'], 1) if hd['posts'] else 0,
        })

    # Heatmap
    heatmap = {}
    for p in posts:
        if p['hour'] is not None and 1 <= p['day'] <= 7:
            key = f"{p['day']}_{p['hour']}"
            if key not in heatmap:
                heatmap[key] = {'posts': 0, 'eng': 0}
            heatmap[key]['posts'] += 1
            heatmap[key]['eng'] += p['slc']

    D['heatmap'] = []
    for day in range(1, 8):
        for hour in range(24):
            k = f"{day}_{hour}"
            hd = heatmap.get(k, {'posts': 0, 'eng': 0})
            D['heatmap'].append({
                'day': day, 'hour': hour, 'posts': hd['posts'],
                'avg_eng': round(hd['eng'] / hd['posts'], 1) if hd['posts'] else 0,
            })

    # Formats
    fmt_data = defaultdict(lambda: {'count': 0, 'eng': 0})
    for p in posts:
        fmt_data[p['type']]['count'] += 1
        fmt_data[p['type']]['eng'] += p['slc']

    D['formats'] = []
    for t in sorted(fmt_data, key=lambda x: -fmt_data[x]['count']):
        fd = fmt_data[t]
        D['formats'].append({
            'type': t,
            'count': fd['count'],
            'pct': round(fd['count'] / n * 100, 1),
            'avg_eng': round(fd['eng'] / fd['count'], 1) if fd['count'] else 0,
        })

    # Orientation
    ori_map = {'Vertik.': 'Vertikāli', 'Horiz.': 'Horizontāli', 'Kvadrāts': 'Kvadrāts'}
    ori_data = defaultdict(lambda: {'count': 0, 'eng': 0})
    for p in posts:
        if p['a_prop']:
            ori_name = ori_map.get(p['a_prop'], p['a_prop'])
            ori_data[ori_name]['count'] += 1
            ori_data[ori_name]['eng'] += p['slc']

    ori_total = sum(v['count'] for v in ori_data.values())
    D['orientation'] = []
    for name in ['Horizontāli', 'Vertikāli', 'Kvadrāts']:
        od = ori_data[name]
        D['orientation'].append({
            'name': name,
            'count': od['count'],
            'pct': round(od['count'] / ori_total * 100, 1) if ori_total else 0,
            'avg_eng': round(od['eng'] / od['count'], 1) if od['count'] else 0,
        })

    # Caption length
    cap_buckets = [('Nav teksta', 0, 0), ('1-120', 1, 120), ('121-240', 121, 240), ('241+', 241, 99999)]
    cap_data = {b[0]: {'count': 0, 'eng': 0} for b in cap_buckets}
    for p in posts:
        tz = p['t_zimes']
        for name, lo, hi in cap_buckets:
            if lo <= tz <= hi:
                cap_data[name]['count'] += 1
                cap_data[name]['eng'] += p['slc']
                break

    D['caption_length'] = []
    for name, _, _ in cap_buckets:
        cd = cap_data[name]
        D['caption_length'].append({
            'name': name,
            'count': cd['count'],
            'pct': round(cd['count'] / n * 100, 1),
            'avg_eng': round(cd['eng'] / cd['count'], 1) if cd['count'] else 0,
        })

    # Engagement distribution
    eng_buckets = [('0', 0, 0), ('1-50', 1, 50), ('51-100', 51, 100), ('101-200', 101, 200), ('201-500', 201, 500), ('500+', 501, 999999)]
    D['eng_dist'] = []
    for name, lo, hi in eng_buckets:
        cnt = len([p for p in posts if lo <= p['slc'] <= hi])
        D['eng_dist'].append({'bucket': name, 'count': cnt, 'pct': round(cnt / n * 100, 1)})

    # Contest
    contest = [p for p in posts if p['is_contest']]
    non_contest = [p for p in posts if not p['is_contest']]
    D['contest'] = {
        'count': len(contest),
        'pct': round(len(contest) / n * 100, 1),
        'avg_eng': round(sum(p['slc'] for p in contest) / len(contest), 1) if contest else 0,
        'non_avg': round(sum(p['slc'] for p in non_contest) / len(non_contest), 1) if non_contest else 0,
    }

    # Themes
    topic_defs = {
        'Sports': ['sports','futbol','hokej','volejbol','basketbol','sacensīb','čempion','turnīr','spēl','komand','uzvara','medaļ','olimp','atlēt','trenin','stadion'],
        'Kultūra': ['koncert','izstād','teātr','māksla','muzej','galerij','festiv','mūzik','dziesm','dzied','dej','kor','orķestr','kultūr','filmu','kino','grāmat','bibliotēk'],
        'Izglītība': ['skol','student','mācīb','izglītīb','universitet','augstskol','diplom','absolvents','pedagog','skolotāj','lekcij','stipendij','studij'],
        'Pasākumi': ['pasākum','svētk','svinēš','ielūdz','aicin','pieteik','reģistrēj','biļet','ieeja','programm','noris','notik','plkst','apmeklē','piedāvā'],
        'Infrastruktūra': ['celtniecīb','remonts','iela','tilts','parks','laukums','ēka','projekts','būvniecīb','atjauno','labiekārto','infrastruktūr'],
        'Daba/Vide': ['dab','jūr','pludmal','mežs','ezers','ziedi','stādī','apzaļumo','saule','sniegs','rudens','pavasari','vasara','ziem','dabas'],
        'Uzņēmējdarbība': ['uzņēm','biznes','investīcij','darba','vakance','attīstīb','ekonomik','tūrist','viesnīc','restorān','kafejnīc'],
        'Sociālais': ['brīvprātīg','labdarīb','palīdzīb','atbalst','kopiena','iedzīvotāj','ģimene','bērn','jauniet','senior','veselīb'],
    }
    topic_stats = {}
    for topic, keywords in topic_defs.items():
        topic_posts = [p for p in posts if any(kw in p['teksts'].lower() for kw in keywords)]
        if topic_posts:
            topic_stats[topic] = {
                'name': topic,
                'count': len(topic_posts),
                'pct': round(len(topic_posts) / n * 100, 1),
                'avg_eng': round(sum(pp['slc'] for pp in topic_posts) / len(topic_posts), 1),
            }

    D['themes'] = sorted(topic_stats.values(), key=lambda x: -x['count'])

    # Bigrams
    stop_words = {'un','ir','kas','ar','par','no','uz','ka','lai','bet','ja','vai','pie','pa',
                  'šo','šī','tam','to','tā','tas','ko','kā','nav','arī','jau','vēl','tikai',
                  'var','būs','bija','gan','kur','jūs','mēs','viņi','savu','sev','tur','te',
                  'pēc','līdz','bez','caur','starp','priekš','dēļ','laikā','kopš',
                  'aicinām','vairāk','informācija','https','www','com','facebook',
                  'ventspils','ventspilī','ventspilnieku'}
    bigram_counter = Counter()
    for p in posts:
        words = re.findall(r'[a-zA-ZāčēģīķļņšūžĀČĒĢĪĶĻŅŠŪŽ]{3,}', p['teksts'].lower())
        filtered = [w for w in words if w not in stop_words and len(w) > 2]
        for i in range(len(filtered) - 1):
            bigram_counter[filtered[i] + ' ' + filtered[i+1]] += 1

    D['bigrams'] = [{'phrase': bg, 'count': c} for bg, c in bigram_counter.most_common(10)]

    # Top posts
    sorted_top = sorted(posts, key=lambda x: -x['slc'])
    D['top_posts'] = []
    for p in sorted_top[:20]:
        D['top_posts'].append({
            'page': p['page'],
            'type': p['type'],
            'likes': p['likes'],
            'shares': p['shares'],
            'comments': p['comments'],
            'engagement': p['slc'],
            'teksts': p['teksts'][:120],
            'date': p['date'],
            'url': p['fb_url'],
        })

    D['worst_posts'] = []
    for p in sorted_top[-5:]:
        D['worst_posts'].append({
            'page': p['page'],
            'type': p['type'],
            'engagement': p['slc'],
            'teksts': p['teksts'][:80],
            'date': p['date'],
        })

    # Video
    video_posts = [p for p in posts if p['type'] in ('Video', 'Reel') and p['video_len'] > 0]
    D['video'] = {'has_data': len(video_posts) > 0, 'total': len(video_posts)}
    if video_posts:
        vid_buckets = [('0-15s', 0, 15), ('16-30s', 16, 30), ('31-60s', 31, 60), ('61-120s', 61, 120), ('120s+', 121, 99999)]
        vid_data = {b[0]: {'count': 0, 'eng': 0} for b in vid_buckets}
        for p in video_posts:
            for name, lo, hi in vid_buckets:
                if lo <= p['video_len'] <= hi:
                    vid_data[name]['count'] += 1
                    vid_data[name]['eng'] += p['slc']
                    break

        D['video']['buckets'] = []
        for name, _, _ in vid_buckets:
            vd = vid_data[name]
            D['video']['buckets'].append({
                'name': name,
                'count': vd['count'],
                'avg_eng': round(vd['eng'] / vd['count'], 1) if vd['count'] else 0,
            })

    return D


def compute_group(page_posts, research):
    """Compute group-level analytics combining all pages."""
    D = {}
    all_posts = []
    D['pages'] = {}

    for pname in PAGES:
        posts = page_posts[pname]
        all_posts.extend(posts)
        D['pages'][pname] = compute_page(posts, pname)
        print(f"  Computed {pname}: {len(posts)} posts")

    # Group combined KPI
    n = len(all_posts)
    total_likes = sum(p['likes'] for p in all_posts)
    total_shares = sum(p['shares'] for p in all_posts)
    total_comments = sum(p['comments'] for p in all_posts)
    total_eng = total_likes + total_shares + total_comments

    D['group_kpi'] = {
        'name': GROUP_NAME,
        'total_pages': len(PAGES),
        'total_posts': n,
        'total_engagement': total_eng,
        'total_likes': total_likes,
        'total_shares': total_shares,
        'total_comments': total_comments,
        'avg_eng': round(total_eng / n, 1) if n else 0,
    }

    # Research data integration
    D['research'] = {}
    komm = research.get('komunikacija', {})
    if komm:
        D['research']['followers'] = {}
        for p in PAGES:
            D['research']['followers'][p] = safe_int(komm.get('Sekotāju skaits', {}).get(p, 0))
        D['research']['avg_likes'] = {}
        D['research']['avg_comments'] = {}
        D['research']['avg_shares'] = {}
        for p in PAGES:
            D['research']['avg_likes'][p] = komm.get('Vidēji "likes" ziņai', {}).get(p, '0')
            D['research']['avg_comments'][p] = komm.get('Vidēji "comments" ziņai', {}).get(p, '0')
            D['research']['avg_shares'][p] = komm.get('Vidēji "shares" ziņai', {}).get(p, '0')

    pub = research.get('publikacijas', {})
    if pub:
        D['research']['posts_per_day'] = {}
        D['research']['workday_pct'] = {}
        D['research']['weekend_pct'] = {}
        for p in PAGES:
            D['research']['posts_per_day'][p] = pub.get('Vidējais zinu skaits dienā', {}).get(p, '0')
            D['research']['workday_pct'][p] = pub.get('Ziņas darba dienā', {}).get(p, '0')
            D['research']['weekend_pct'][p] = pub.get('Ziņas brīvdienās', {}).get(p, '0')

    # Popular links
    D['saites'] = research.get('saites', [])

    # Combined top posts (from all pages together)
    all_sorted = sorted(all_posts, key=lambda x: -x['slc'])
    D['top_posts'] = []
    for p in all_sorted[:20]:
        D['top_posts'].append({
            'page': p['page'],
            'type': p['type'],
            'likes': p['likes'],
            'shares': p['shares'],
            'comments': p['comments'],
            'engagement': p['slc'],
            'teksts': p['teksts'][:120],
            'date': p['date'],
            'url': p['fb_url'],
        })

    D['worst_posts'] = []
    for p in all_sorted[-10:]:
        D['worst_posts'].append({
            'page': p['page'],
            'type': p['type'],
            'engagement': p['slc'],
            'teksts': p['teksts'][:80],
            'date': p['date'],
        })

    return D


def generate_html(D):
    data_json = json.dumps(D, ensure_ascii=False)
    pages_json = json.dumps(PAGES, ensure_ascii=False)
    short_json = json.dumps(PAGE_SHORT, ensure_ascii=False)
    colors_json = json.dumps(PAGE_COLORS, ensure_ascii=False)
    num_pages = len(PAGES)

    section_num = [0]
    def sn():
        section_num[0] += 1
        return f'{section_num[0]:02d}'

    s = {}
    s['compare'] = sn()
    s['monthly_eng'] = sn()
    s['monthly_avg'] = sn()
    s['formats'] = sn()
    s['fmt_eff'] = sn()
    s['orientation'] = sn()
    s['caption'] = sn()
    s['weekday'] = sn()
    s['hourly'] = sn()
    s['work_wknd'] = sn()
    s['heatmap'] = sn()
    s['eng_dist'] = sn()
    s['h1h2'] = sn()
    s['contest'] = sn()
    s['themes'] = sn()
    s['bigrams'] = sn()
    s['saites'] = sn()
    s['top_posts'] = sn()
    s['worst'] = sn()

    # Build dynamic HTML parts
    legend_html = ''.join(
        f'<span class="legend-item"><span class="legend-dot" style="background:{PAGE_COLORS[p]}"></span>{PAGE_SHORT.get(p,p)}</span>'
        for p in PAGES
    )
    heatmap_html = ''.join(
        f'<div><h4 style="color:{PAGE_COLORS[p]};margin-bottom:8px">{PAGE_SHORT.get(p,p)}</h4><div class="hm-grid" id="hm{i}"></div></div>'
        for i, p in enumerate(PAGES)
    )
    bigram_html = ''.join(f'<div id="bg{i}"></div>' for i in range(num_pages))
    hm_cols = min(num_pages, 3)
    bg_cols = min(num_pages, 3)

    # ── Build the JavaScript as a raw string (no f-string escaping needed) ──
    js_code = r'''
const D=__DATA_JSON__;
const PAGES=__PAGES_JSON__;
const SHORT=__SHORT_JSON__;
const COLORS=__COLORS_JSON__;
const PD=PAGES.map(function(p){return D.pages[p]});

function nf(n){return n.toString()}
function sn(p){return SHORT[p]||p}

Chart.defaults.color='#8B95A5';
Chart.defaults.borderColor='rgba(255,255,255,0.04)';
Chart.defaults.font.family="'Inter',system-ui,sans-serif";
Chart.defaults.font.size=12;
Chart.defaults.plugins.legend.labels.padding=16;
Chart.defaults.plugins.legend.labels.usePointStyle=true;
Chart.defaults.plugins.legend.labels.pointStyle='circle';
Chart.defaults.plugins.tooltip.backgroundColor='#1A2030';
Chart.defaults.plugins.tooltip.borderColor='rgba(255,255,255,0.08)';
Chart.defaults.plugins.tooltip.borderWidth=1;
Chart.defaults.plugins.tooltip.cornerRadius=10;
Chart.defaults.plugins.tooltip.padding=12;
Chart.defaults.plugins.tooltip.titleFont={weight:'600',size:13};
Chart.defaults.elements.bar.borderRadius=8;
Chart.defaults.elements.bar.borderSkipped=false;
Chart.defaults.elements.line.borderWidth=2;
Chart.defaults.elements.point.radius=3;
Chart.defaults.elements.point.hoverRadius=6;

function scOpts(){return{grid:{color:'rgba(255,255,255,0.04)'},ticks:{color:'#8B95A5',font:{size:11}}}}
var CH={};
function makeChart(cid,cfg){var el=document.getElementById(cid);if(!el)return;if(CH[cid])CH[cid].destroy();CH[cid]=new Chart(el,cfg)}
function alpha(hex,a){var r=parseInt(hex.slice(1,3),16),g=parseInt(hex.slice(3,5),16),b=parseInt(hex.slice(5,7),16);return 'rgba('+r+','+g+','+b+','+a+')'}

// ── KPI ──
function buildKPI(){
  var g=D.group_kpi;
  var postParts=PAGES.map(function(p){return sn(p)+': '+D.pages[p].kpi.total_posts}).join(' · ');
  var engParts=PAGES.map(function(p){return sn(p)+': '+nf(D.pages[p].kpi.total_engagement)}).join(' · ');
  var avgParts=PAGES.map(function(p){return sn(p)+': '+D.pages[p].kpi.avg_eng}).join(' · ');
  var items=[
    ['Lapas grupā',g.total_pages,'__GROUP_NAME__'],
    ['Publikācijas',nf(g.total_posts),postParts],
    ['Kopējā iesaiste',nf(g.total_engagement),engParts],
    ['Vid. iesaiste',g.avg_eng,avgParts],
    ['Likes',nf(g.total_likes),'Shares: '+nf(g.total_shares)+' · Comments: '+nf(g.total_comments)],
  ];
  document.getElementById('kpiG').innerHTML=items.map(function(it){
    return '<div class="kpi"><div class="label">'+it[0]+'</div><div class="value">'+it[1]+'</div><div class="sub">'+it[2]+'</div></div>';
  }).join('');
}

// ── COMPARE TABLE ──
function buildCompare(){
  var r=D.research||{};
  var fol=r.followers||{};
  var ppd=r.posts_per_day||{};
  var wkd=r.workday_pct||{};
  var wke=r.weekend_pct||{};

  var metrics=[
    {label:'Sekotāji',fn:function(p,d){return fol[p]||'-'},higher:true,format:'nf'},
    {label:'Publikācijas',fn:function(p,d){return d.kpi.total_posts},higher:true},
    {label:'Kopējā iesaiste',fn:function(p,d){return d.kpi.total_engagement},higher:true,format:'nf'},
    {label:'Vid. iesaiste',fn:function(p,d){return d.kpi.avg_eng},higher:true},
    {label:'Mediāna',fn:function(p,d){return d.kpi.median_eng},higher:true},
    {label:'Virālie posti',fn:function(p,d){return d.kpi.viral_count+' ('+d.kpi.viral_pct+'%)'},higher:true,raw:true},
    {label:'Posti/dienā',fn:function(p,d){return ppd[p]||'-'},higher:false,raw:true},
    {label:'Darba dienās',fn:function(p,d){return wkd[p]||'-'},higher:false,raw:true},
    {label:'Brīvdienās',fn:function(p,d){return wke[p]||'-'},higher:false,raw:true},
  ];

  var h='<table class="compare-table"><tr><th></th>';
  PAGES.forEach(function(p){h+='<th style="color:'+COLORS[p]+';font-weight:700">'+sn(p)+'</th>'});
  h+='</tr>';

  metrics.forEach(function(m){
    h+='<tr><td class="cm-label">'+m.label+'</td>';
    var vals=PAGES.map(function(p){return {p:p,v:m.fn(p,D.pages[p])}});
    // Find best value for winner badge
    var numVals=vals.map(function(x){return parseFloat(String(x.v).replace(/[^0-9.]/g,''))});
    var bestIdx=-1;
    if(!m.raw){
      var best=m.higher?Math.max.apply(null,numVals):Math.min.apply(null,numVals);
      bestIdx=numVals.indexOf(best);
    }
    vals.forEach(function(x,i){
      var v=m.format==='nf'?nf(x.v):x.v;
      var badge=i===bestIdx?' <span class="winner winner-yes">★</span>':'';
      h+='<td style="font-weight:600;color:'+COLORS[x.p]+'">'+v+badge+'</td>';
    });
    h+='</tr>';
  });
  h+='</table>';
  document.getElementById('compareG').innerHTML=h;
}

// ── MONTHLY ENGAGEMENT ──
function buildMonthlyEng(){
  var labels=PD[0].monthly.map(function(x){return x.name});
  var datasets=PAGES.map(function(p,i){
    return {label:sn(p),data:D.pages[p].monthly.map(function(x){return x.engagement}),backgroundColor:alpha(COLORS[p],0.4),borderColor:COLORS[p],borderWidth:2};
  });
  makeChart('mEngC',{type:'bar',data:{labels:labels,datasets:datasets},options:{responsive:true,maintainAspectRatio:false,interaction:{mode:'index',intersect:false},scales:{x:scOpts(),y:scOpts()}}});
}

// ── MONTHLY AVG ──
function buildMonthlyAvg(){
  var labels=PD[0].monthly.map(function(x){return x.name});
  var datasets=PAGES.map(function(p,i){
    return {label:sn(p),data:D.pages[p].monthly.map(function(x){return x.avg_eng}),borderColor:COLORS[p],backgroundColor:alpha(COLORS[p],0.06),fill:true,borderWidth:3,pointRadius:4,pointBackgroundColor:COLORS[p],tension:.4};
  });
  makeChart('mAvgC',{type:'line',data:{labels:labels,datasets:datasets},options:{responsive:true,maintainAspectRatio:false,interaction:{mode:'index',intersect:false},scales:{x:scOpts(),y:scOpts()}}});
}

// ── FORMATS ──
function buildFormats(){
  var allTypes={};
  PAGES.forEach(function(p){D.pages[p].formats.forEach(function(f){allTypes[f.type]=true})});
  var types=Object.keys(allTypes);
  function getVal(fmts,t,key){var f=fmts.find(function(x){return x.type===t});return f?f[key]:0}
  var ds1=PAGES.map(function(p){return {label:sn(p),data:types.map(function(t){return getVal(D.pages[p].formats,t,'count')}),backgroundColor:alpha(COLORS[p],0.5),borderColor:COLORS[p],borderWidth:2}});
  makeChart('fmtC',{type:'bar',data:{labels:types,datasets:ds1},options:{responsive:true,maintainAspectRatio:false,scales:{x:scOpts(),y:scOpts()}}});
  var ds2=PAGES.map(function(p){return {label:sn(p),data:types.map(function(t){return getVal(D.pages[p].formats,t,'avg_eng')}),backgroundColor:alpha(COLORS[p],0.5),borderColor:COLORS[p],borderWidth:2}});
  makeChart('fmtEffC',{type:'bar',data:{labels:types,datasets:ds2},options:{responsive:true,maintainAspectRatio:false,indexAxis:'y',scales:{x:scOpts(),y:scOpts()}}});
}

// ── ORIENTATION ──
function buildOri(){
  var names=['Horizontāli','Vertikāli','Kvadrāts'];
  function getAvg(oris,n){var o=oris.find(function(x){return x.name===n});return o?o.avg_eng:0}
  var ds=PAGES.map(function(p){return {label:sn(p),data:names.map(function(n){return getAvg(D.pages[p].orientation,n)}),backgroundColor:alpha(COLORS[p],0.5),borderColor:COLORS[p],borderWidth:2}});
  makeChart('oriC',{type:'bar',data:{labels:names,datasets:ds},options:{responsive:true,maintainAspectRatio:false,scales:{x:scOpts(),y:scOpts()}}});
}

// ── CAPTION LENGTH ──
function buildCaption(){
  var names=PD[0].caption_length.map(function(x){return x.name});
  function getAvg(caps,n){var c=caps.find(function(x){return x.name===n});return c?c.avg_eng:0}
  var ds=PAGES.map(function(p){return {label:sn(p),data:names.map(function(n){return getAvg(D.pages[p].caption_length,n)}),backgroundColor:alpha(COLORS[p],0.5),borderColor:COLORS[p],borderWidth:2}});
  makeChart('capC',{type:'bar',data:{labels:names,datasets:ds},options:{responsive:true,maintainAspectRatio:false,scales:{x:scOpts(),y:scOpts()}}});
}

// ── WEEKDAY ──
function buildWeekday(){
  var labels=PD[0].weekday.map(function(x){return x.name});
  var ds=PAGES.map(function(p){return {label:sn(p),data:D.pages[p].weekday.map(function(x){return x.avg_eng}),borderColor:COLORS[p],backgroundColor:alpha(COLORS[p],0.06),fill:false,borderWidth:3,pointRadius:4,pointBackgroundColor:COLORS[p],tension:.4}});
  makeChart('wdayC',{type:'line',data:{labels:labels,datasets:ds},options:{responsive:true,maintainAspectRatio:false,interaction:{mode:'index',intersect:false},scales:{x:scOpts(),y:scOpts()}}});
}

// ── HOURLY ──
function buildHourly(){
  var labels=PD[0].hourly.map(function(x){return x.hour});
  var ds=PAGES.map(function(p){return {label:sn(p),data:D.pages[p].hourly.map(function(x){return x.avg_eng}),borderColor:COLORS[p],backgroundColor:'transparent',fill:false,borderWidth:2,pointRadius:2,pointBackgroundColor:COLORS[p],tension:.4}});
  makeChart('hourC',{type:'line',data:{labels:labels,datasets:ds},options:{responsive:true,maintainAspectRatio:false,interaction:{mode:'index',intersect:false},scales:{x:scOpts(),y:scOpts()}}});
}

// ── WORK vs WEEKEND ──
function buildWW(){
  var ds=PAGES.map(function(p){return {label:sn(p),data:[D.pages[p].work_weekend.work_avg,D.pages[p].work_weekend.weekend_avg],backgroundColor:[alpha(COLORS[p],0.5),alpha(COLORS[p],0.3)],borderColor:COLORS[p],borderWidth:2}});
  makeChart('wwC',{type:'bar',data:{labels:['Darba dienas','Brīvdienas'],datasets:ds},options:{responsive:true,maintainAspectRatio:false,scales:{x:scOpts(),y:scOpts()}}});
}

// ── HEATMAPS ──
function buildHeatmap(hm,elId){
  var days=['','P','O','T','C','Pk','S','Sv'];
  var vals=hm.filter(function(x){return x.avg_eng>0}).map(function(x){return x.avg_eng});
  vals.sort(function(a,b){return b-a});
  var top5=vals.length>=5?vals[4]:vals[vals.length-1]||0;
  var bot5vals=vals.slice().sort(function(a,b){return a-b});
  var bot5=bot5vals.length>=5?bot5vals[4]:bot5vals[bot5vals.length-1]||0;
  var html='<div class="hm-label"></div>';
  for(var h=0;h<24;h++) html+='<div class="hm-label">'+h+'</div>';
  for(var d=1;d<=7;d++){
    html+='<div class="hm-label">'+days[d]+'</div>';
    for(var h=0;h<24;h++){
      var cell=hm.find(function(x){return x.day===d&&x.hour===h});
      var val=cell?cell.avg_eng:0;
      var bg;
      if(val<=0){bg='rgba(255,255,255,0.02)'}
      else if(val>=top5){bg='rgba(52,211,153,0.75)'}
      else if(val<=bot5){bg='rgba(239,68,68,0.65)'}
      else{var mid=(val-bot5)/(top5-bot5);bg='rgba(251,191,36,'+(0.2+mid*0.5)+')'}
      html+='<div class="hm-cell" style="background:'+bg+'" title="'+days[d]+' '+h+':00 — vid. '+val+'">'+(val>0?Math.round(val):'')+'</div>';
    }
  }
  document.getElementById(elId).innerHTML=html;
}

// ── ENGAGEMENT DIST ──
function buildEngDist(){
  var labels=PD[0].eng_dist.map(function(x){return x.bucket});
  var ds=PAGES.map(function(p){return {label:sn(p),data:D.pages[p].eng_dist.map(function(x){return x.pct}),backgroundColor:alpha(COLORS[p],0.5),borderColor:COLORS[p],borderWidth:2}});
  makeChart('engDistC',{type:'bar',data:{labels:labels,datasets:ds},options:{responsive:true,maintainAspectRatio:false,scales:{x:scOpts(),y:Object.assign(scOpts(),{title:{display:true,text:'% no publikācijām',color:'#8B95A5',font:{size:11}}})}}});
}

// ── H1 vs H2 ──
function buildH1H2(){
  var ds=PAGES.map(function(p){return {label:sn(p),data:[D.pages[p].half_year[0].avg_eng,D.pages[p].half_year[1].avg_eng],backgroundColor:[alpha(COLORS[p],0.5),alpha(COLORS[p],0.3)],borderColor:COLORS[p],borderWidth:2}});
  makeChart('halfC',{type:'bar',data:{labels:['H1 (Jan-Jun)','H2 (Jul-Dec)'],datasets:ds},options:{responsive:true,maintainAspectRatio:false,scales:{x:scOpts(),y:scOpts()}}});
}

// ── CONTESTS ──
function buildContest(){
  var ds=PAGES.map(function(p){var c=D.pages[p].contest;return {label:sn(p),data:[c.avg_eng,c.non_avg],backgroundColor:[alpha(COLORS[p],0.5),alpha(COLORS[p],0.3)],borderColor:COLORS[p],borderWidth:2}});
  makeChart('contestC',{type:'bar',data:{labels:['Konkursi vid.','Parastie vid.'],datasets:ds},options:{responsive:true,maintainAspectRatio:false,scales:{x:scOpts(),y:scOpts()}}});
}

// ── THEMES ──
function buildThemes(){
  var allTopics={};
  PAGES.forEach(function(p){D.pages[p].themes.forEach(function(t){allTopics[t.name]=true})});
  var topics=Object.keys(allTopics);
  function getAvg(themes,n){var t=themes.find(function(x){return x.name===n});return t?t.avg_eng:0}
  var ds=PAGES.map(function(p){return {label:sn(p),data:topics.map(function(n){return getAvg(D.pages[p].themes,n)}),backgroundColor:alpha(COLORS[p],0.5),borderColor:COLORS[p],borderWidth:2}});
  makeChart('themesC',{type:'bar',data:{labels:topics,datasets:ds},options:{responsive:true,maintainAspectRatio:false,indexAxis:'y',scales:{x:scOpts(),y:scOpts()}}});
}

// ── BIGRAMS ──
function buildBigrams(){
  PAGES.forEach(function(p,i){
    var bg=D.pages[p].bigrams;
    if(!bg||!bg.length)return;
    var maxC=bg[0].count;
    var c=COLORS[p];
    var html='<h4 style="color:'+c+';font-size:0.95em;margin-bottom:10px;font-family:Space Grotesk,sans-serif">'+sn(p)+'</h4>';
    html+='<table><tr><th>#</th><th>Vārdu pāris</th><th style="text-align:right">Reizes</th><th style="width:35%"></th></tr>';
    bg.forEach(function(b,j){
      var pct=Math.round(b.count/maxC*100);
      html+='<tr><td>'+(j+1)+'</td><td style="font-weight:600;color:var(--text-1)">'+b.phrase+'</td><td style="text-align:right;color:'+c+'">'+b.count+'</td><td><div style="background:'+c+'33;height:14px;border-radius:4px;width:'+pct+'%"></div></td></tr>';
    });
    html+='</table>';
    document.getElementById('bg'+i).innerHTML=html;
  });
}

// ── SAITES ──
function buildSaites(){
  var s=D.saites;
  if(!s||!s.length)return;
  var html='<table><tr><th>#</th><th>Domēns</th><th style="text-align:right">Reizes</th><th style="text-align:right">%</th></tr>';
  s.forEach(function(x,i){
    html+='<tr><td>'+(i+1)+'</td><td style="font-weight:600;color:var(--text-1)">'+x.domain+'</td><td style="text-align:right;color:var(--accent)">'+x.count+'</td><td style="text-align:right;color:var(--text-2)">'+x.pct+'</td></tr>';
  });
  html+='</table>';
  document.getElementById('saitesG').innerHTML=html;
}

// ── TOP POSTS ──
var topSortKey='engagement';
function buildTopPosts(){
  var t=D.top_posts.slice().sort(function(a,b){return b[topSortKey]-a[topSortKey]});
  var h='<tr><th>#</th><th>Lapa</th><th>Datums</th><th>Tips</th><th>Likes</th><th>Shares</th><th>Comm.</th><th>Kopā</th><th>Teksts</th></tr>';
  t.forEach(function(p,i){
    var c=COLORS[p.page]||'#8B95A5';
    h+='<tr><td>'+(i+1)+'</td><td><span class="b" style="background:'+c+'22;color:'+c+'">'+sn(p.page)+'</span></td><td>'+p.date+'</td><td><span class="b bb">'+p.type+'</span></td><td>'+nf(p.likes)+'</td><td>'+nf(p.shares)+'</td><td>'+nf(p.comments)+'</td><td><strong>'+nf(p.engagement)+'</strong></td><td style="max-width:200px;font-size:0.8em;color:var(--text-3)">'+(p.url?'<a href="'+p.url+'" target="_blank">':'')+p.teksts.substring(0,60)+'…'+(p.url?'</a>':'')+'</td></tr>';
  });
  document.getElementById('topT').innerHTML=h;
}
function sortTop(key,btn){topSortKey=key;btn.parentElement.querySelectorAll('button').forEach(function(b){b.classList.remove('active')});btn.classList.add('active');buildTopPosts()}

function buildWorstPosts(){
  var t=D.worst_posts;
  var h='<tr><th>#</th><th>Lapa</th><th>Datums</th><th>Tips</th><th>Iesaiste</th><th>Teksts</th></tr>';
  t.forEach(function(p,i){
    var c=COLORS[p.page]||'#8B95A5';
    h+='<tr><td>'+(i+1)+'</td><td><span class="b" style="background:'+c+'22;color:'+c+'">'+sn(p.page)+'</span></td><td>'+p.date+'</td><td><span class="b br">'+p.type+'</span></td><td>'+p.engagement+'</td><td style="max-width:220px;font-size:0.8em;color:var(--text-3)">'+p.teksts.substring(0,50)+'…</td></tr>';
  });
  document.getElementById('worstT').innerHTML=h;
}

// ── CONCLUSIONS ──
function buildConclusions(){
  var items=[];
  // Find best/worst by avg engagement
  var sorted=PAGES.slice().sort(function(a,b){return D.pages[b].kpi.avg_eng-D.pages[a].kpi.avg_eng});
  var best=sorted[0],worst=sorted[sorted.length-1];

  items.push(['good','Lielākā vid. iesaiste: '+sn(best),'<span style="color:'+COLORS[best]+'">'+sn(best)+'</span> sasniedz augstāko vidējo iesaisti: <span class="st">'+D.pages[best].kpi.avg_eng+'</span>. Zemākā: <span style="color:'+COLORS[worst]+'">'+sn(worst)+'</span> ar <span class="st">'+D.pages[worst].kpi.avg_eng+'</span>.']);

  // Most posts
  var mostPosts=PAGES.slice().sort(function(a,b){return D.pages[b].kpi.total_posts-D.pages[a].kpi.total_posts})[0];
  items.push(['tip','Visvairāk publikāciju: '+sn(mostPosts),'<span style="color:'+COLORS[mostPosts]+'">'+sn(mostPosts)+'</span> publicē visvairāk — <span class="st">'+D.pages[mostPosts].kpi.total_posts+'</span> ierakstu gadā.']);

  // Viral comparison
  var mostViral=PAGES.slice().sort(function(a,b){return D.pages[b].kpi.viral_pct-D.pages[a].kpi.viral_pct})[0];
  items.push(['good','Virālākais saturs: '+sn(mostViral),'<span style="color:'+COLORS[mostViral]+'">'+sn(mostViral)+'</span> ģenerē visvairāk virālo saturu — <span class="st">'+D.pages[mostViral].kpi.viral_pct+'%</span> no visiem postiem.']);

  // Best day per page
  var dayInfo=PAGES.map(function(p){var bd=D.pages[p].weekday.reduce(function(a,b){return a.avg_eng>b.avg_eng?a:b});return '<span style="color:'+COLORS[p]+'">'+sn(p)+'</span>: '+bd.name+' ('+bd.avg_eng+')'}).join(', ');
  items.push(['tip','Labākās dienas',dayInfo]);

  // Best format per page
  var fmtInfo=PAGES.map(function(p){var bf=D.pages[p].formats.reduce(function(a,b){return a.avg_eng>b.avg_eng?a:b});return '<span style="color:'+COLORS[p]+'">'+sn(p)+'</span>: '+bf.type+' ('+bf.avg_eng+')'}).join(', ');
  items.push(['tip','Efektīvākais formāts',fmtInfo]);

  // Text length
  var capInfo=PAGES.map(function(p){var bc=D.pages[p].caption_length.reduce(function(a,b){return a.avg_eng>b.avg_eng?a:b});return '<span style="color:'+COLORS[p]+'">'+sn(p)+'</span>: '+bc.name+' ('+bc.avg_eng+')'}).join(', ');
  items.push(['tip','Labākais teksta garums',capInfo]);

  // Contest impact
  var contestInfo=PAGES.map(function(p){var c=D.pages[p].contest;return '<span style="color:'+COLORS[p]+'">'+sn(p)+'</span>: konkursi '+c.avg_eng+' vs parastie '+c.non_avg}).join(', ');
  items.push(['good','Konkursu efekts',contestInfo]);

  // Theme insights
  var themeInfo=PAGES.map(function(p){if(!D.pages[p].themes.length)return '';var bt=D.pages[p].themes.slice().sort(function(a,b){return b.avg_eng-a.avg_eng})[0];return '<span style="color:'+COLORS[p]+'">'+sn(p)+'</span>: '+bt.name+' ('+bt.avg_eng+')'}).filter(function(x){return x}).join(', ');
  if(themeInfo)items.push(['good','Efektīvākā tēma',themeInfo]);

  // Action plan
  items.push(['tip','Ieteikumi grupai','<p>1. Dalīties ar labāko praksi — katrai lapai ir savas stiprās puses.</p><p>2. Optimizēt publicēšanas laiku uz labākajām dienām un stundām.</p><p>3. Testēt efektīvākos formātus no citām grupas lapām.</p><p>4. Palielināt brīvdienu saturu visās lapās.</p>']);

  var g=document.getElementById('conG');
  g.innerHTML=items.map(function(it){
    return '<div class="con '+it[0]+'"><h3>'+it[1]+'</h3><p>'+it[2]+'</p></div>';
  }).join('');
}

function initScroll(){
  var obs=new IntersectionObserver(function(e){e.forEach(function(en){if(en.isIntersecting)en.target.classList.add('visible')})},{threshold:0.08,rootMargin:'0px 0px -40px 0px'});
  document.querySelectorAll('.section,.con').forEach(function(el){obs.observe(el)});
}

document.addEventListener('DOMContentLoaded',function(){
  buildKPI();
  buildCompare();
  buildMonthlyEng();
  buildMonthlyAvg();
  buildFormats();
  buildOri();
  buildCaption();
  buildWeekday();
  buildHourly();
  buildWW();
  PAGES.forEach(function(p,i){buildHeatmap(D.pages[p].heatmap,'hm'+i)});
  buildEngDist();
  buildH1H2();
  buildContest();
  buildThemes();
  buildBigrams();
  buildSaites();
  buildTopPosts();
  buildWorstPosts();
  buildConclusions();
  document.querySelectorAll('.con').forEach(function(el,i){el.style.transitionDelay=(i*0.06)+'s'});
  initScroll();
});
'''
    # Inject data into JS template
    js_code = js_code.replace('__DATA_JSON__', data_json)
    js_code = js_code.replace('__PAGES_JSON__', pages_json)
    js_code = js_code.replace('__SHORT_JSON__', short_json)
    js_code = js_code.replace('__COLORS_JSON__', colors_json)
    js_code = js_code.replace('__GROUP_NAME__', GROUP_NAME)

    # ── Build dynamic HTML fragments ──
    c_first = PAGE_COLORS[PAGES[0]]
    c_second = PAGE_COLORS[PAGES[1 % len(PAGES)]]

    # ── Prev/Next group navigation ──
    idx = GROUP_ORDER.index(GROUP_KEY) if GROUP_KEY in GROUP_ORDER else 0
    prev_key = GROUP_ORDER[idx - 1] if idx > 0 else None
    next_key = GROUP_ORDER[idx + 1] if idx < len(GROUP_ORDER) - 1 else None

    nav_pn_parts = []
    if prev_key:
        pg = GROUPS[prev_key]
        nav_pn_parts.append(f'<a href="Facebook_{pg["slug"]}_2025.html">&larr; {pg["name"]}</a>')
    nav_pn_parts.append(f'<span class="acc-name">{GROUP_NAME}</span>')
    if next_key:
        ng = GROUPS[next_key]
        nav_pn_parts.append(f'<a href="Facebook_{ng["slug"]}_2025.html">{ng["name"]} &rarr;</a>')
    nav_pn_html = '\n'.join(nav_pn_parts)

    badge_parts = ' &middot; '.join(PAGE_SHORT.get(p,p) for p in PAGES)

    # ── Unified footer with all group links ──
    uf_colors = ['#1877F2', '#34D399', '#FBBF24', '#F87171', '#A78BFA', '#06b6d4', '#ec4899', '#f97316']
    uf_links = []
    uf_links.append('<a class="uf-link summary-link" href="Facebook_apak&#353;v&#299;tra_visas_2025.html" title="Kop&#275;jais p&#257;rskats"><span class="uf-dot" style="background:linear-gradient(135deg,#1877F2,#A78BFA)">K</span>Kop&#275;jais p&#257;rskats</a>')
    for gi, gk in enumerate(GROUP_ORDER):
        g = GROUPS[gk]
        active = ' active' if gk == GROUP_KEY else ''
        letter = g['name'][0]
        color = uf_colors[gi % len(uf_colors)]
        uf_links.append(f'<a class="uf-link{active}" href="Facebook_{g["slug"]}_2025.html" title="{g["name"]}"><span class="uf-dot" style="background:{color}">{letter}</span>{g["name"]}</a>')
    unified_footer_html = '\n'.join(uf_links)

    hm_per_row = min(num_pages, 2)
    bg_per_row = min(num_pages, 3)

    heatmap_sections_html = ''
    for i, p in enumerate(PAGES):
        heatmap_sections_html += f'<div><h4 style="color:{PAGE_COLORS[p]};margin-bottom:8px;font-family:Space Grotesk,sans-serif;font-size:0.95em">{PAGE_SHORT.get(p,p)}</h4><div class="hm-grid" id="hm{i}"></div></div>\n'

    bigram_divs_html = ''.join(f'<div id="bg{i}"></div>' for i in range(num_pages))

    desc_all = f'Visu {num_pages} lapu'

    # ── Build HTML using raw string template (no brace escaping needed) ──
    html_template = r'''<!DOCTYPE html>
<html lang="lv"><head><meta charset="UTF-8"><meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>__GROUP_NAME__ — Facebook 2025</title>
<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800;900&family=Space+Grotesk:wght@500;600;700&display=swap" rel="stylesheet">
<script src="https://cdnjs.cloudflare.com/ajax/libs/Chart.js/4.4.1/chart.umd.min.js"></script>
<style>
:root {
  --bg: #060911;
  --surface-1: #131820;
  --surface-2: #1A2030;
  --surface-3: #232B3B;
  --border: rgba(255,255,255,0.05);
  --border-hover: rgba(255,255,255,0.12);
  --text-1: #F0F2F5;
  --text-2: #8B95A5;
  --text-2b: #A8B2C1;
  --text-3: #5A6577;
  --accent: #1877F2;
  --accent-hover: #4293F5;
  --accent-dim: rgba(24,119,242,0.12);
  --green: #34D399;
  --green-dim: rgba(52,211,153,0.12);
  --red: #F87171;
  --red-dim: rgba(248,113,113,0.12);
  --orange: #FBBF24;
  --orange-dim: rgba(251,191,36,0.12);
  --purple: #A78BFA;
  --purple-dim: rgba(167,139,250,0.12);
  --chart-grid: rgba(255,255,255,0.04);
  --radius-card: 20px;
  --radius-btn: 12px;
}
*{margin:0;padding:0;box-sizing:border-box}
html{scroll-behavior:smooth}
body{
  font-family:'Inter',system-ui,sans-serif;
  background:var(--bg);
  background-image:
    radial-gradient(ellipse 80% 60% at 20% 10%, rgba(24,119,242,0.06) 0%, transparent 60%),
    radial-gradient(ellipse 60% 50% at 80% 80%, rgba(52,211,153,0.05) 0%, transparent 60%);
  background-attachment:fixed;
  color:var(--text-1);line-height:1.6;-webkit-font-smoothing:antialiased;
}
body::after{
  content:'';position:fixed;top:0;left:0;width:100%;height:100%;
  opacity:0.015;pointer-events:none;z-index:9999;
  background-image:url("data:image/svg+xml,%3Csvg viewBox='0 0 256 256' xmlns='http://www.w3.org/2000/svg'%3E%3Cfilter id='n'%3E%3CfeTurbulence type='fractalNoise' baseFrequency='0.9' numOctaves='4' stitchTiles='stitch'/%3E%3C/filter%3E%3Crect width='100%25' height='100%25' filter='url(%23n)'/%3E%3C/svg%3E");
}
.nav{position:sticky;top:0;z-index:100;background:rgba(6,9,17,0.85);backdrop-filter:blur(16px);padding:14px 48px;display:flex;align-items:center;justify-content:space-between;gap:12px;border-bottom:1px solid var(--border);font-size:0.82em;flex-wrap:wrap}
.nav a{color:var(--accent);text-decoration:none;font-weight:600;display:flex;align-items:center;gap:6px;transition:color .2s;white-space:nowrap}.nav a:hover{color:var(--accent-hover)}
.nav .acc-name{color:var(--text-1);font-family:'Space Grotesk',sans-serif;font-size:1.05em;font-weight:700}
.nav-pn{display:flex;gap:16px;align-items:center}
.unified-footer{margin-top:120px;padding:64px 48px 40px;background:rgba(10,13,20,0.8);backdrop-filter:blur(20px);border-top:1px solid rgba(255,255,255,0.04)}
.unified-footer .uf-title{font-size:1.3em;margin-bottom:32px;text-align:center;background:linear-gradient(135deg,var(--text-2),var(--text-3));-webkit-background-clip:text;-webkit-text-fill-color:transparent}
.unified-footer .uf-grid{display:flex;flex-wrap:wrap;justify-content:center;gap:10px 12px;max-width:1280px;margin:0 auto 32px}
.unified-footer .uf-link{font-size:0.78em;padding:8px 16px 8px 10px;border-radius:16px;background:rgba(19,24,32,0.6);backdrop-filter:blur(8px);border:1px solid rgba(255,255,255,0.04);color:var(--text-2b);text-decoration:none;display:inline-flex;align-items:center;gap:8px;font-weight:500;transition:all .2s ease}
.unified-footer .uf-link:hover{background:rgba(35,43,59,0.8);border-color:rgba(79,140,255,0.15);box-shadow:0 4px 20px rgba(0,0,0,0.3);transform:translateY(-3px);color:var(--text-1)}
.unified-footer .uf-link.active{border-color:rgba(24,119,242,0.3);background:rgba(24,119,242,0.08);color:var(--text-1)}
.unified-footer .uf-dot{width:22px;height:22px;border-radius:50%;display:inline-flex;align-items:center;justify-content:center;font-size:0.72em;font-weight:700;color:#fff;flex-shrink:0}
.unified-footer .uf-copy{font-size:0.75em;color:var(--text-3);text-align:center;padding-top:24px;max-width:1280px;margin:0 auto}
.container{max-width:1400px;margin:0 auto;padding:0 48px 120px}
.hero{padding:80px 48px 60px;text-align:center;position:relative;overflow:hidden}
.hero::before{content:'';position:absolute;top:0;left:50%;transform:translateX(-50%);width:900px;height:600px;background:radial-gradient(ellipse,rgba(24,119,242,0.12) 0%,rgba(52,211,153,0.08) 40%,transparent 70%);animation:heroGlow 8s ease-in-out infinite alternate;pointer-events:none}
@keyframes heroGlow{0%{opacity:0.8;transform:translateX(-50%) scale(1)}100%{opacity:1;transform:translateX(-50%) scale(1.15)}}
.hero h1{font-family:'Space Grotesk',sans-serif;font-size:2.8em;font-weight:700;letter-spacing:-0.04em;line-height:1.1;margin-bottom:12px}
.hero h1 span{background:linear-gradient(135deg,__C_FIRST__,__C_SECOND__);-webkit-background-clip:text;-webkit-text-fill-color:transparent}
.subtitle{color:var(--text-2b);font-size:1.05em;opacity:0.7}
.group-badge{display:inline-block;padding:6px 20px;border-radius:20px;background:linear-gradient(135deg,rgba(24,119,242,0.15),rgba(52,211,153,0.15));border:1px solid rgba(255,255,255,0.08);font-size:0.85em;font-weight:600;color:var(--text-2b);margin-top:16px}

/* KPI grid */
.kpi-grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(160px,1fr));gap:18px;margin-bottom:60px}
.kpi{background:rgba(19,24,32,0.6);backdrop-filter:blur(16px);border-radius:20px;padding:22px 18px;border:1px solid rgba(255,255,255,0.05);box-shadow:0 4px 24px rgba(0,0,0,0.2),inset 0 1px 0 rgba(255,255,255,0.03);text-align:center;transition:all .3s ease;animation:fadeInUp 0.6s ease-out both}
.kpi:hover{border-color:rgba(24,119,242,0.15);transform:translateY(-4px)}
.kpi .label{font-size:0.72em;letter-spacing:2px;color:var(--text-3);text-transform:uppercase;font-weight:600}
.kpi .value{font-size:1.65em;font-weight:800;margin:8px 0;background:linear-gradient(135deg,var(--text-1),var(--text-2b));-webkit-background-clip:text;-webkit-text-fill-color:transparent;font-variant-numeric:tabular-nums;white-space:nowrap}
.kpi .sub{font-size:0.7em;color:var(--text-2)}
@keyframes fadeInUp{from{opacity:0;transform:translateY(20px)}to{opacity:1;transform:translateY(0)}}
.kpi:nth-child(1){animation-delay:.05s}.kpi:nth-child(2){animation-delay:.1s}.kpi:nth-child(3){animation-delay:.15s}.kpi:nth-child(4){animation-delay:.2s}.kpi:nth-child(5){animation-delay:.25s}.kpi:nth-child(6){animation-delay:.3s}

/* Compare table */
.compare-table{width:100%;border-collapse:collapse;font-size:0.85em}
.compare-table th{padding:14px 16px;text-align:center;font-weight:700;font-size:0.85em;border-bottom:2px solid rgba(24,119,242,0.15)}
.compare-table td{padding:12px 16px;text-align:center;border-bottom:1px solid var(--border);font-variant-numeric:tabular-nums}
.compare-table .cm-label{text-align:left;color:var(--text-2);font-weight:500;font-size:0.9em}
.compare-table tr:hover td{background:rgba(24,119,242,0.03)}

/* Section styles */
.stitle{font-family:'Space Grotesk',sans-serif;font-size:1.7em;font-weight:700;color:var(--text-1);letter-spacing:-0.02em;margin-bottom:16px;padding-left:4px}
.stitle .num{font-size:0.6em;color:rgba(24,119,242,0.4);margin-right:10px;font-weight:700}
.sdesc{color:var(--text-2b);font-size:0.9em;line-height:1.65;max-width:700px;margin-bottom:16px;padding-left:4px}
.section{opacity:0;transform:translateY(40px);transition:opacity .9s cubic-bezier(.16,1,.3,1),transform .9s cubic-bezier(.16,1,.3,1);margin-bottom:80px}
.section.visible{opacity:1;transform:translateY(0)}
.cs{background:rgba(19,24,32,0.5);backdrop-filter:blur(12px);border-radius:20px;padding:28px;border:1px solid rgba(255,255,255,0.04);box-shadow:0 2px 16px rgba(0,0,0,0.15),inset 0 1px 0 rgba(255,255,255,0.02);transition:all .3s ease}
.cs:hover{border-color:rgba(255,255,255,0.08)}
.cw{position:relative;height:380px}.cw.tall{height:520px}
.two-col{display:grid;grid-template-columns:1fr 1fr;gap:28px}
.hm-wrap{display:grid;grid-template-columns:repeat(__HM_COLS__,1fr);gap:28px}
.bg-wrap{display:grid;grid-template-columns:repeat(__BG_COLS__,1fr);gap:28px}
table{width:100%;border-collapse:collapse;font-size:0.82em;margin-top:8px}
th,td{padding:12px 14px;text-align:left;border-bottom:1px solid var(--border)}
th{color:var(--text-3);font-weight:600;font-size:0.72em;text-transform:uppercase;letter-spacing:1px;border-bottom:2px solid rgba(24,119,242,0.1)}
td{font-variant-numeric:tabular-nums;color:var(--text-2b)}
tr:hover td{background:rgba(24,119,242,0.03)}
a{color:var(--accent);text-decoration:none}a:hover{color:var(--accent-hover)}
.b{display:inline-block;padding:3px 12px;border-radius:12px;font-size:0.74em;font-weight:600}
.bg{background:var(--green-dim);color:var(--green)}.br{background:var(--red-dim);color:var(--red)}.bb{background:var(--accent-dim);color:var(--accent)}.bo{background:var(--orange-dim);color:var(--orange)}
.sd2{font-family:'Space Grotesk',sans-serif;text-align:center;font-size:1.4em;font-weight:700;color:var(--text-2);letter-spacing:-0.01em;margin:80px 0 50px;padding:20px 0;position:relative}
.sd2::after{content:'';position:absolute;bottom:0;left:50%;transform:translateX(-50%);width:80px;height:3px;border-radius:2px;background:linear-gradient(90deg,__C_FIRST__,__C_SECOND__)}
.con-grid{display:grid;grid-template-columns:1fr 1fr;gap:24px;max-width:1400px;margin:0 auto}
.con{background:rgba(19,24,32,0.5);backdrop-filter:blur(12px);border-radius:20px;padding:28px;border:1px solid rgba(255,255,255,0.04);border-left:4px solid var(--accent);box-shadow:0 2px 16px rgba(0,0,0,0.15);transition:all .5s cubic-bezier(.16,1,.3,1);opacity:0;transform:translateY(30px)}
.con.visible{opacity:1;transform:translateY(0)}
.con:hover{transform:translateY(-3px)}
.con h3{margin-bottom:12px;color:var(--accent);font-size:1.1em;font-weight:700}
.con p{color:var(--text-2b);font-size:0.88em;margin-bottom:8px;line-height:1.7}
.hl{color:var(--text-1);font-weight:600}.st{color:var(--orange);font-weight:700}
.good{border-left-color:var(--green)}.good h3{color:var(--green)}
.bad{border-left-color:var(--red)}.bad h3{color:var(--red)}
.tip{border-left-color:var(--orange)}.tip h3{color:var(--orange)}
.hm-grid{display:grid;grid-template-columns:40px repeat(24,1fr);gap:2px;font-size:0.65em}
.hm-cell{border-radius:4px;text-align:center;padding:4px 2px;font-weight:500;min-height:28px;display:flex;align-items:center;justify-content:center}
.hm-label{color:var(--text-3);font-weight:600;display:flex;align-items:center;justify-content:center;font-size:0.9em}
.cc button{padding:5px 14px;border-radius:12px;cursor:pointer;font-size:0.74em;font-weight:500;font-family:'Inter',sans-serif;transition:all .2s;border:1px solid var(--border);background:var(--surface-2);color:var(--text-2)}
.cc button.active{background:var(--accent);color:#fff;border-color:var(--accent)}
.cc button:hover:not(.active){border-color:var(--border-hover);color:var(--text-1)}
.ch{display:flex;justify-content:flex-end;align-items:center;gap:8px;margin-bottom:12px}
.legend-inline{display:flex;gap:20px;margin-bottom:12px;padding-left:4px;flex-wrap:wrap}
.legend-dot{width:12px;height:12px;border-radius:50%;display:inline-block;margin-right:6px;vertical-align:middle}
.legend-item{font-size:0.82em;color:var(--text-2b);font-weight:500}
.winner{display:inline-block;padding:2px 8px;border-radius:8px;font-size:0.65em;font-weight:700;margin-left:8px;vertical-align:middle}
.winner-yes{background:var(--green-dim);color:var(--green)}
::-webkit-scrollbar{width:8px}::-webkit-scrollbar-thumb{background:rgba(24,119,242,0.15);border-radius:4px}
::selection{background:rgba(24,119,242,0.25);color:var(--text-1)}
@media(max-width:900px){.two-col,.con-grid,.hm-wrap,.bg-wrap{grid-template-columns:1fr}.container{padding:0 20px 80px}.hero{padding:60px 20px 40px}.hero h1{font-size:2em}.nav{padding:10px 16px;font-size:0.75em}.unified-footer{padding:40px 20px 30px}.unified-footer .uf-grid{gap:8px}}
</style>
</head><body>

<div class="nav">
<a href="Facebook_apak&#353;v&#299;tra_visas_2025.html">&larr; Kop&#275;jais</a>
<div class="nav-pn">
__NAV_PN__
</div>
</div>

<div class="hero">
<h1><span>__GROUP_NAME__</span></h1>
<p class="subtitle">Facebook 2025 &middot; Grupu sal&#299;dzin&#257;jums &middot; Janv&#257;ris &#8212; Decembris</p>
<div class="group-badge">__BADGE_TEXT__</div>
</div>

<div class="container">
<div class="kpi-grid" id="kpiG"></div>

<div class="legend-inline">__LEGEND_HTML__</div>

<div class="section"><h2 class="stitle"><span class="num">__S_COMPARE__</span>Lapu sal&#299;dzin&#257;jums</h2>
<p class="sdesc">__DESC_ALL__ galven&#257;s metrikas &#8212; &#9733; apz&#299;m&#275; lab&#257;ko rezult&#257;tu.</p>
<div class="cs" id="compareG"></div></div>

<div class="section"><h2 class="stitle"><span class="num">__S_MONTHLY_ENG__</span>Kop&#275;j&#257; iesaiste pa m&#275;ne&#353;iem</h2>
<p class="sdesc">__DESC_ALL__ kop&#275;j&#257; iesaiste (likes+shares+comments) katr&#257; m&#275;nes&#299;.</p>
<div class="cs"><div class="cw" style="height:420px"><canvas id="mEngC"></canvas></div></div></div>

<div class="section"><h2 class="stitle"><span class="num">__S_MONTHLY_AVG__</span>Vid&#275;j&#257; iesaiste pa m&#275;ne&#353;iem</h2>
<p class="sdesc">Vid&#275;j&#257; iesaiste uz vienu publik&#257;ciju katr&#257; m&#275;nes&#299;.</p>
<div class="cs"><div class="cw" style="height:380px"><canvas id="mAvgC"></canvas></div></div></div>

<div class="section"><div class="two-col">
<div><h2 class="stitle"><span class="num">__S_FORMATS__</span>Form&#257;tu sadal&#299;jums</h2>
<p class="sdesc">Publik&#257;ciju skaits pa form&#257;tiem.</p>
<div class="cs"><div class="cw"><canvas id="fmtC"></canvas></div></div></div>
<div><h2 class="stitle"><span class="num">__S_FMT_EFF__</span>Form&#257;tu efektivit&#257;te</h2>
<p class="sdesc">Vid&#275;j&#257; iesaiste p&#275;c form&#257;ta.</p>
<div class="cs"><div class="cw"><canvas id="fmtEffC"></canvas></div></div></div>
</div></div>

<div class="section"><div class="two-col">
<div><h2 class="stitle"><span class="num">__S_ORIENTATION__</span>Bil&#382;u orient&#257;cija</h2>
<div class="cs"><div class="cw"><canvas id="oriC"></canvas></div></div></div>
<div><h2 class="stitle"><span class="num">__S_CAPTION__</span>Teksta garums</h2>
<p class="sdesc">Vid&#275;j&#257; iesaiste p&#275;c pavado&#353;&#257; teksta garuma.</p>
<div class="cs"><div class="cw"><canvas id="capC"></canvas></div></div></div>
</div></div>

<div class="section"><h2 class="stitle"><span class="num">__S_WEEKDAY__</span>Iesaiste p&#275;c ned&#275;&#316;as dien&#257;m</h2>
<div class="cs"><div class="cw"><canvas id="wdayC"></canvas></div></div></div>

<div class="section"><h2 class="stitle"><span class="num">__S_HOURLY__</span>Iesaiste pa stund&#257;m</h2>
<div class="cs"><div class="cw tall"><canvas id="hourC"></canvas></div></div></div>

<div class="section"><h2 class="stitle"><span class="num">__S_WORK_WKND__</span>Darba dienas vs Br&#299;vdienas</h2>
<div class="cs"><div class="cw"><canvas id="wwC"></canvas></div></div></div>

<div class="section"><h2 class="stitle"><span class="num">__S_HEATMAP__</span>Karstuma kartes</h2>
<p class="sdesc">Vid&#275;j&#257; iesaiste pa ned&#275;&#316;as dien&#257;m un stund&#257;m &#8212; za&#316;&#353; = lab&#257;kais, sarkans = slikt&#257;kais.</p>
<div class="cs"><div class="hm-wrap">__HEATMAP_HTML__</div></div></div>

<div class="sd2">Papildu Anal&#299;ze</div>

<div class="section"><h2 class="stitle"><span class="num">__S_ENG_DIST__</span>Iesaistes sadal&#299;jums</h2>
<p class="sdesc">Cik publik&#257;ciju iekr&#299;t katr&#257; iesaistes diapazon&#257;.</p>
<div class="cs"><div class="cw"><canvas id="engDistC"></canvas></div></div></div>

<div class="section"><div class="two-col">
<div><h2 class="stitle"><span class="num">__S_H1H2__</span>H1 vs H2 sal&#299;dzin&#257;jums</h2>
<div class="cs"><div class="cw"><canvas id="halfC"></canvas></div></div></div>
<div><h2 class="stitle"><span class="num">__S_CONTEST__</span>Konkursi vs Parastie</h2>
<div class="cs"><div class="cw"><canvas id="contestC"></canvas></div></div></div>
</div></div>

<div class="section"><h2 class="stitle"><span class="num">__S_THEMES__</span>Satura t&#275;mas</h2>
<p class="sdesc">T&#275;mu klasifik&#257;cija un vid&#275;j&#257; iesaiste.</p>
<div class="cs"><div class="cw tall"><canvas id="themesC"></canvas></div></div></div>

<div class="section"><h2 class="stitle"><span class="num">__S_BIGRAMS__</span>Bie&#382;&#257;kie v&#257;rdu p&#257;ri</h2>
<p class="sdesc">Top 10 bie&#382;&#257;k lietotie 2 v&#257;rdu p&#257;ri katr&#257; lap&#257;.</p>
<div class="cs"><div class="bg-wrap">__BIGRAM_HTML__</div></div></div>

<div class="section"><h2 class="stitle"><span class="num">__S_SAITES__</span>Popul&#257;r&#257;k&#257;s saites</h2>
<p class="sdesc">Bie&#382;&#257;k izmantotie dom&#275;ni satur&#257;.</p>
<div class="cs"><div id="saitesG"></div></div></div>

<div class="sd2">Top Ieraksti</div>

<div class="section"><h2 class="stitle"><span class="num">__S_TOP_POSTS__</span>Top 20 publik&#257;cijas</h2>
<div class="cs"><div class="ch"><div class="cc"><button class="active" onclick="sortTop('engagement',this)">Iesaiste</button><button onclick="sortTop('likes',this)">Likes</button><button onclick="sortTop('shares',this)">Shares</button><button onclick="sortTop('comments',this)">Comments</button></div></div><div style="overflow-x:auto"><table id="topT"></table></div></div></div>

<div class="section"><h2 class="stitle"><span class="num">__S_WORST__</span>Zem&#257;k&#257; iesaiste</h2>
<div class="cs"><div style="overflow-x:auto"><table id="worstT"></table></div></div></div>

<div class="sd2">Secin&#257;jumi un Ieteikumi</div>
<div class="con-grid" id="conG"></div>

</div>

<div class="unified-footer">
<div class="uf-title">Visi Facebook grupas p&#257;rskati 2025</div>
<div class="uf-grid">
__UNIFIED_FOOTER__
</div>
<div class="uf-copy">Ventspils Facebook 2025 &mdash; 8 grupas &middot; 31 lapa &middot; Grupu p&#257;rskati</div>
</div>

<script>
__JS_CODE__
</script>
</body></html>'''

    # Inject all placeholders
    html = html_template
    html = html.replace('__JS_CODE__', js_code)
    html = html.replace('__GROUP_NAME__', GROUP_NAME)
    html = html.replace('__NAV_PN__', nav_pn_html)
    html = html.replace('__BADGE_TEXT__', badge_parts)
    html = html.replace('__LEGEND_HTML__', legend_html)
    html = html.replace('__HEATMAP_HTML__', heatmap_sections_html)
    html = html.replace('__BIGRAM_HTML__', bigram_divs_html)
    html = html.replace('__UNIFIED_FOOTER__', unified_footer_html)
    html = html.replace('__C_FIRST__', c_first)
    html = html.replace('__C_SECOND__', c_second)
    html = html.replace('__DESC_ALL__', desc_all)
    html = html.replace('__HM_COLS__', str(hm_per_row))
    html = html.replace('__BG_COLS__', str(bg_per_row))
    html = html.replace('__S_COMPARE__', s['compare'])
    html = html.replace('__S_MONTHLY_ENG__', s['monthly_eng'])
    html = html.replace('__S_MONTHLY_AVG__', s['monthly_avg'])
    html = html.replace('__S_FORMATS__', s['formats'])
    html = html.replace('__S_FMT_EFF__', s['fmt_eff'])
    html = html.replace('__S_ORIENTATION__', s['orientation'])
    html = html.replace('__S_CAPTION__', s['caption'])
    html = html.replace('__S_WEEKDAY__', s['weekday'])
    html = html.replace('__S_HOURLY__', s['hourly'])
    html = html.replace('__S_WORK_WKND__', s['work_wknd'])
    html = html.replace('__S_HEATMAP__', s['heatmap'])
    html = html.replace('__S_ENG_DIST__', s['eng_dist'])
    html = html.replace('__S_H1H2__', s['h1h2'])
    html = html.replace('__S_CONTEST__', s['contest'])
    html = html.replace('__S_THEMES__', s['themes'])
    html = html.replace('__S_BIGRAMS__', s['bigrams'])
    html = html.replace('__S_SAITES__', s['saites'])
    html = html.replace('__S_TOP_POSTS__', s['top_posts'])
    html = html.replace('__S_WORST__', s['worst'])

    return html


def main():
    page_posts = load_all_data()
    research = load_research()
    D = compute_group(page_posts, research)

    html = generate_html(D)

    out_path = os.path.join(DIR, f'Facebook_{GROUP_SLUG}_2025.html')
    with open(out_path, 'w', encoding='utf-8') as f:
        f.write(html)

    file_size = os.path.getsize(out_path) // 1024
    print(f"\nGenerated: {out_path}")
    print(f"File size: {file_size} KB")
    print(f"Group: {GROUP_NAME} ({len(PAGES)} pages)")


if __name__ == '__main__':
    main()
